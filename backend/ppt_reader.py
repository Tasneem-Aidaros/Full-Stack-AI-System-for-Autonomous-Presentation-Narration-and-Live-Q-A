from pptx import Presentation


def read_presentation(file_path):
    presentation = Presentation(file_path)

    slides_data = []

    for slide_number, slide in enumerate(presentation.slides, start=1):

        slide_text = ""

        for shape in slide.shapes:

            if shape.has_text_frame:
                slide_text += shape.text + "\n"

        slides_data.append({
            "slide": slide_number,
            "text": slide_text.strip(),
            "image": f"slide_images/Slide{slide_number}.PNG"
        })

    return slides_data